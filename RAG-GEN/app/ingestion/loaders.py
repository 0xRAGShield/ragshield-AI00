from __future__ import annotations

import hashlib
import logging
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.models.evidence_models import (
    ImageDocument,
    SourceDocument,
)


logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class ParserLimits:
    """Safety limits for ingestion file parsing."""

    max_file_size_bytes: int = 100 * 1024 * 1024
    max_pdf_pages: int = 500
    max_extracted_text_chars: int = 10_000_000
    max_directory_files: int = 10_000


class FileParser:
    """
    Parse supported source files into canonical ingestion documents.

    Responsibilities:
    - Validate file paths and supported formats.
    - Extract textual content from supported documents.
    - Represent images as image documents.
    - Preserve basic source provenance.
    - Enforce parser safety limits.

    Forbidden responsibilities:
    - OCR / computer vision.
    - Cleaning / normalization.
    - Chunking.
    - Embedding generation.
    - Vector indexing.
    - Retrieval / reranking.
    - LLM generation.
    - Business logic.
    """

    DOCUMENT_EXTENSIONS = frozenset(
        {
            ".pdf",
            ".txt",
            ".md",
            ".csv",
            ".docx",
        }
    )

    IMAGE_EXTENSIONS = frozenset(
        {
            ".jpg",
            ".jpeg",
            ".png",
            ".webp",
            ".bmp",
            ".tif",
            ".tiff",
        }
    )

    SUPPORTED_EXTENSIONS = (
        DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS
    )

    _MEDIA_TYPES = {
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".csv": "text/csv",
        ".pdf": "application/pdf",
        ".docx": (
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
    }

    _IMAGE_MEDIA_TYPES = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
        ".bmp": "image/bmp",
        ".tif": "image/tiff",
        ".tiff": "image/tiff",
    }

    def __init__(
        self,
        limits: ParserLimits | None = None,
    ) -> None:
        self._limits = limits or ParserLimits()

        self._validate_limits()

    def load_file(
        self,
        file_path: str | Path,
    ) -> SourceDocument | ImageDocument:
        """
        Parse one supported file.

        Raises:
            FileNotFoundError:
                File does not exist.
            ValueError:
                Path is invalid, unsupported, empty, or violates limits.
            OSError:
                File cannot be accessed.
        """

        path = Path(file_path)

        self._validate_file(path)

        extension = path.suffix.lower()

        if extension in self.DOCUMENT_EXTENSIONS:
            return self._parse_document(path)

        return self._parse_image(path)

    def load_directory(
        self,
        directory_path: str | Path,
    ) -> list[SourceDocument | ImageDocument]:
        """
        Recursively load supported files from a directory.

        Invalid individual files are skipped and logged so one bad
        source does not abort the entire ingestion batch.
        """

        directory = Path(directory_path)

        if not directory.exists():
            raise FileNotFoundError(
                f"Directory not found: {directory}"
            )

        if not directory.is_dir():
            raise ValueError(
                f"Path is not a directory: {directory}"
            )

        documents: list[
            SourceDocument | ImageDocument
        ] = []

        processed_files = 0

        for file_path in sorted(directory.rglob("*")):
            if not file_path.is_file():
                continue

            if file_path.suffix.lower() not in self.SUPPORTED_EXTENSIONS:
                continue

            processed_files += 1

            if processed_files > self._limits.max_directory_files:
                raise ValueError(
                    "Directory contains more supported files than "
                    f"the configured limit "
                    f"({self._limits.max_directory_files})"
                )

            try:
                document = self.load_file(file_path)
                documents.append(document)

            except (OSError, ValueError) as error:
                logger.warning(
                    "Skipping ingestion file",
                    extra={
                        "file_path": str(file_path),
                        "error_type": type(error).__name__,
                        "error": str(error),
                    },
                )

        return documents

    def _validate_file(self, file_path: Path) -> None:
        if not file_path.exists():
            raise FileNotFoundError(
                f"File not found: {file_path}"
            )

        if not file_path.is_file():
            raise ValueError(
                f"Path is not a file: {file_path}"
            )

        extension = file_path.suffix.lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type: {extension or '<none>'}"
            )

        try:
            file_size = file_path.stat().st_size
        except OSError:
            raise

        if file_size <= 0:
            raise ValueError(
                f"Empty file: {file_path.name}"
            )

        if file_size > self._limits.max_file_size_bytes:
            raise ValueError(
                f"File exceeds maximum allowed size: "
                f"{file_path.name}"
            )

    def _parse_document(
        self,
        file_path: Path,
    ) -> SourceDocument:
        extension = file_path.suffix.lower()

        if extension == ".pdf":
            return self._parse_pdf(file_path)

        if extension == ".docx":
            return self._parse_docx(file_path)

        text = self._read_text_file(file_path)

        return self._build_source_document(
            file_path=file_path,
            text=text,
        )

    def _parse_pdf(
        self,
        file_path: Path,
    ) -> SourceDocument:
        reader = PdfReader(str(file_path))

        page_count = len(reader.pages)

        if page_count > self._limits.max_pdf_pages:
            raise ValueError(
                f"PDF exceeds maximum page count: "
                f"{file_path.name}"
            )

        pages: list[dict[str, int | str]] = []
        total_chars = 0

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):
            text = page.extract_text() or ""
            text = text.strip()

            if not text:
                continue

            total_chars += len(text)

            if (
                total_chars
                > self._limits.max_extracted_text_chars
            ):
                raise ValueError(
                    f"Extracted PDF text exceeds maximum allowed "
                    f"size: {file_path.name}"
                )

            pages.append(
                {
                    "page": page_number,
                    "text": text,
                }
            )

        if not pages:
            raise ValueError(
                f"No extractable text found in PDF: "
                f"{file_path.name}"
            )

        full_text = "\n\n".join(
            str(page["text"])
            for page in pages
        )

        return SourceDocument(
            source_id=self._build_source_id(file_path),
            title=file_path.name,
            text=full_text,
            media_type="application/pdf",
            metadata={
                "pages": pages,
                "page_count": page_count,
                "extracted_page_count": len(pages),
                "content_sha256": self._compute_sha256(file_path),
            },
        )

    def _parse_docx(
        self,
        file_path: Path,
    ) -> SourceDocument:
        document = Document(str(file_path))

        paragraphs: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        full_text = "\n\n".join(paragraphs)

        if not full_text:
            raise ValueError(
                f"No extractable text found in DOCX: "
                f"{file_path.name}"
            )

        self._validate_extracted_text(full_text, file_path)

        return SourceDocument(
            source_id=self._build_source_id(file_path),
            title=file_path.name,
            text=full_text,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            metadata={
                "content_sha256": self._compute_sha256(file_path),
            },
        )

    def _parse_image(
        self,
        file_path: Path,
    ) -> ImageDocument:
        extension = file_path.suffix.lower()

        return ImageDocument(
            source_id=self._build_source_id(file_path),
            title=file_path.name,
            image_path=str(file_path),
            media_type=self._get_image_media_type(extension),
            metadata={
                "file_type": "image",
                "extension": extension,
                "content_sha256": self._compute_sha256(file_path),
            },
        )

    def _read_text_file(
        self,
        file_path: Path,
    ) -> str:
        try:
            text = file_path.read_text(
                encoding="utf-8",
            )
        except UnicodeDecodeError as error:
            raise ValueError(
                f"File is not valid UTF-8 text: "
                f"{file_path.name}"
            ) from error

        text = text.strip()

        if not text:
            raise ValueError(
                f"Empty file: {file_path.name}"
            )

        self._validate_extracted_text(
            text,
            file_path,
        )

        return text

    def _validate_extracted_text(
        self,
        text: str,
        file_path: Path,
    ) -> None:
        if len(text) > self._limits.max_extracted_text_chars:
            raise ValueError(
                f"Extracted text exceeds maximum allowed size: "
                f"{file_path.name}"
            )

    def _build_source_id(
        self,
        file_path: Path,
    ) -> str:
        """
        Build a stable source identifier from the canonical path.

        The identifier is intentionally deterministic for the same
        ingestion source path.
        """

        try:
            canonical_path = file_path.resolve()
        except OSError:
            canonical_path = file_path.absolute()

        digest = hashlib.sha256(
            str(canonical_path).encode("utf-8")
        ).hexdigest()

        return digest

    def _compute_sha256(
        self,
        file_path: Path,
    ) -> str:
        digest = hashlib.sha256()

        with file_path.open("rb") as file:
            for chunk in iter(
                lambda: file.read(1024 * 1024),
                b"",
            ):
                digest.update(chunk)

        return digest.hexdigest()

    def _get_image_media_type(
        self,
        extension: str,
    ) -> str:
        media_type = self._IMAGE_MEDIA_TYPES.get(extension)

        if media_type is None:
            raise ValueError(
                f"Unsupported image extension: {extension}"
            )

        return media_type

    def _validate_limits(self) -> None:
        if self._limits.max_file_size_bytes <= 0:
            raise ValueError(
                "max_file_size_bytes must be greater than zero"
            )

        if self._limits.max_pdf_pages <= 0:
            raise ValueError(
                "max_pdf_pages must be greater than zero"
            )

        if self._limits.max_extracted_text_chars <= 0:
            raise ValueError(
                "max_extracted_text_chars must be greater than zero"
            )

        if self._limits.max_directory_files <= 0:
            raise ValueError(
                "max_directory_files must be greater than zero"
            )