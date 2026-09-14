from pathlib import Path

from pypdf import PdfReader
from docx import Document

from app.models.evidence_models import SourceDocument


class FileParser:

    # Supported document formats
    DOCUMENT_EXTENSIONS = {
        ".pdf",
        ".txt",
        ".md",
        ".csv",
        ".docx"
    }

    # Supported image formats
    IMAGE_EXTENSIONS = {
        ".jpg",
        ".jpeg",
        ".png",
        ".webp",
        ".bmp",
        ".tif",
        ".tiff"
    }

    SUPPORTED_EXTENSIONS = (
        DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS
    )

    def load_file(self, file_path: str):

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"File not found: {file_path}"
            )

        if not path.is_file():
            raise ValueError(
                f"Path is not a file: {file_path}"
            )

        extension = path.suffix.lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type: {extension}. "
                f"Supported types: {self.SUPPORTED_EXTENSIONS}"
            )

        # Documents
        if extension in self.DOCUMENT_EXTENSIONS:
            return self._parse_document(path)

        # Images
        if extension in self.IMAGE_EXTENSIONS:
            return self._parse_image(path)

    def _parse_document(
        self,
        file_path: Path
    ) -> SourceDocument:

        extension = file_path.suffix.lower()

        # PDF
        if extension == ".pdf":
            return self._parse_pdf(file_path)

        # DOCX
        if extension == ".docx":
            return self._parse_docx(file_path)

        # TXT / MD / CSV
        text = file_path.read_text(
            encoding="utf-8"
        ).strip()

        if not text:
            raise ValueError(
                f"Empty file: {file_path.name}"
            )

        return SourceDocument(
            source_id=file_path.stem,
            title=file_path.name,
            text=text,
            media_type=self._get_media_type(extension)
        )

    def _parse_pdf(
        self,
        file_path: Path
    ) -> SourceDocument:

        reader = PdfReader(str(file_path))

        pages = []

        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):
            text = page.extract_text()

            if text and text.strip():
                pages.append({
                    "page": page_number,
                    "text": text.strip()
                })

        if not pages:
            raise ValueError(
                f"No extractable text found in PDF: "
                f"{file_path.name}"
            )

        full_text = "\n\n".join(
            page["text"]
            for page in pages
        )

        return SourceDocument(
            source_id=file_path.stem,
            title=file_path.name,
            text=full_text,
            media_type="application/pdf",
            metadata={
                "pages": pages
            }
        )

    def _parse_docx(
        self,
        file_path: Path
    ) -> SourceDocument:

        document = Document(str(file_path))

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        full_text = "\n\n".join(paragraphs)

        if not full_text:
            raise ValueError(
                f"No extractable text found in DOCX: "
                f"{file_path.name}"
            )

        return SourceDocument(
            source_id=file_path.stem,
            title=file_path.name,
            text=full_text,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            )
        )

    def _parse_image(
        self,
        file_path: Path
    ) -> dict:

        return {
            "type": "image",
            "image_path": str(file_path),
            "metadata": {
                "source": file_path.name,
                "file_path": str(file_path),
                "file_type": "image",
                "extension": file_path.suffix.lower()
            }
        }

    def _get_media_type(
        self,
        extension: str
    ) -> str:

        media_types = {
            ".txt": "text/plain",
            ".md": "text/markdown",
            ".csv": "text/csv"
        }

        return media_types.get(
            extension,
            "text/plain"
        )

    def load_directory(
        self,
        directory_path: str
    ) -> list:

        directory = Path(directory_path)

        if not directory.exists():
            raise FileNotFoundError(
                f"Directory not found: {directory_path}"
            )

        if not directory.is_dir():
            raise ValueError(
                f"Path is not a directory: {directory_path}"
            )

        files = []

        for file_path in directory.rglob("*"):

            if not file_path.is_file():
                continue

            if file_path.suffix.lower() not in self.SUPPORTED_EXTENSIONS:
                continue

            try:
                file = self.load_file(
                    str(file_path)
                )

                files.append(file)

            except ValueError as error:
                print(
                    f"Skipping {file_path.name}: {error}"
                )

        return files